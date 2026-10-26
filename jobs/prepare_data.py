"""Download UCI original, preserve SHA-256 provenance, stream both XLSX sheets to CSV."""
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import urlopen
from zipfile import ZipFile
from openpyxl import load_workbook

URL = "https://archive.ics.uci.edu/static/public/502/online+retail+ii.zip"
ROOT = Path("/app/data")
FIELDS = ["invoice_id", "product_id", "description", "quantity", "invoice_ts",
          "unit_price", "customer_id", "country", "source_sheet", "source_row"]
ALIASES = {"Invoice": "invoice_id", "InvoiceNo": "invoice_id", "StockCode": "product_id",
           "Description": "description", "Quantity": "quantity", "InvoiceDate": "invoice_ts",
           "Price": "unit_price", "UnitPrice": "unit_price", "Customer ID": "customer_id",
           "CustomerID": "customer_id", "Country": "country"}


def main():
    ROOT.mkdir(parents=True, exist_ok=True)
    archive = ROOT / "online_retail_ii.zip"
    if not archive.exists():
        temp = archive.with_suffix(".download")
        with urlopen(URL, timeout=120) as response, temp.open("wb") as out:
            while chunk := response.read(1024 * 1024):
                out.write(chunk)
        temp.replace(archive)
    with ZipFile(archive) as zipped:
        names = [n for n in zipped.namelist() if n.lower().endswith(".xlsx")]
        if len(names) != 1:
            raise ValueError(f"Expected one workbook, got {names}")
        workbook_path = ROOT / "online_retail_II.xlsx"
        with zipped.open(names[0]) as source, workbook_path.open("wb") as out:
            while chunk := source.read(1024 * 1024):
                out.write(chunk)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    target = ROOT / "retail.csv"
    temporary = ROOT / "retail.csv.partial"
    counts = {}
    try:
        with temporary.open("w", newline="", encoding="utf-8") as out:
            writer = csv.DictWriter(out, fieldnames=FIELDS)
            writer.writeheader()
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows)
                mapped = [ALIASES.get(str(h).strip()) for h in headers]
                if set(FIELDS[:8]) - set(mapped):
                    raise ValueError(f"Unexpected columns in {sheet.title}: {headers}")
                counts[sheet.title] = 0
                for index, values in enumerate(rows, start=2):
                    if all(v is None for v in values):
                        continue
                    record = {key: value for key, value in zip(mapped, values) if key}
                    timestamp = record["invoice_ts"]
                    if isinstance(timestamp, datetime):
                        record["invoice_ts"] = timestamp.isoformat(sep=" ")
                    record.update(source_sheet=sheet.title, source_row=index)
                    writer.writerow(record)
                    counts[sheet.title] += 1
        temporary.replace(target)
    finally:
        workbook.close()
    manifest = {"source_url": URL, "dataset_doi": "10.24432/C5CG6D", "license": "CC BY 4.0",
                "prepared_at_utc": datetime.now(timezone.utc).isoformat(),
                "archive_sha256": hashlib.sha256(archive.read_bytes()).hexdigest(),
                "csv_sha256": hashlib.sha256(target.read_bytes()).hexdigest(),
                "rows_by_sheet": counts, "rows": sum(counts.values())}
    (ROOT / "manifest.json").write_text(json.dumps(manifest, indent=2))
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
